"""Add a static, clickable Kindle-compatible table of contents to a DOCX."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def _delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def _bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _hyperlink_paragraph(text: str, anchor: str):
    paragraph = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "100")
    ppr.append(spacing)
    paragraph.append(ppr)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(Pt(11).pt * 2)))
    rpr.extend((color, underline, size))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend((rpr, text_node))
    hyperlink.append(run)
    paragraph.append(hyperlink)
    return paragraph


def add_toc(source: Path, output: Path) -> None:
    document = Document(source)
    paragraphs = document.paragraphs
    contents_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Contents")
    contents = paragraphs[contents_index]

    # Remove the non-clickable generated contents text between the Contents
    # heading and the real Introduction heading.
    for paragraph in list(document.paragraphs[contents_index + 1:]):
        if paragraph.style.name == "Heading 1" and paragraph.text.strip() == "Introduction":
            break
        _delete_paragraph(paragraph)

    targets = []
    seen_contents = False
    for paragraph in document.paragraphs:
        if paragraph._p is contents._p:
            seen_contents = True
            continue
        if not seen_contents or paragraph.style.name != "Heading 1":
            continue
        title = paragraph.text.strip()
        if not title:
            continue
        anchor = f"kdp_toc_{len(targets) + 1}"
        _bookmark(paragraph, anchor, 1000 + len(targets))
        targets.append((title, anchor))

    cursor = contents._p
    for title, anchor in targets:
        node = _hyperlink_paragraph(title, anchor)
        cursor.addnext(node)
        cursor = node

    document.settings.element.set(qn("w:updateFields"), "true")
    document.save(output)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: add_kdp_toc.py INPUT.docx OUTPUT.docx")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if source == output:
        backup = source.with_name(source.stem + ".pre-toc.docx")
        if not backup.exists():
            shutil.copy2(source, backup)
        temp = source.with_name(source.stem + ".toc-temp.docx")
        add_toc(source, temp)
        temp.replace(source)
    else:
        add_toc(source, output)


if __name__ == "__main__":
    main()
